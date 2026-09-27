import copy
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from ocr_roi_validator.automation import TestCase, load_cases, parse_expected_text
from ocr_roi_validator.automation_ocr import prepare_rois, validate_cases
from scripts.create_vesta_tc_template import build_template


class BlockParserTests(unittest.TestCase):
    def test_multiple_blocks_preserve_internal_whitespace_and_normalize_newlines(self):
        text = "\r\n[ROI3]\r\n\r\n  troisième  \r\n\r\n次の行\r\n\r\n[roi1]\rA: B\r[ROI12]\n마지막\n"
        self.assertEqual(parse_expected_text(text), {3: "  troisième  \n\n次の行", 1: "A: B", 12: "마지막"})

    def test_invalid_blocks_fail_instead_of_silently_treating_markers_as_text(self):
        bad = ["", "   ", "roi1:hello", "prefix\n[ROI1]\nhello", "[ROI1]", "[ROI1]\n \n[ROI2]\nb",
               "[ROI1]\na\n[roi1]\nb", "[ROI01]\na\n[ROI1]\nb", "[ROI0]\na", "[ROI-1]\na",
               "[ROI1]\na\n[ROI2]:b", "[ROI1]\na\n[ROI 2]\nb", "[ROI1]\na\n[ROIbad]\nb", 123]
        for text in bad:
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_expected_text(text)

    def test_marker_like_text_inside_sentence_is_not_a_separator(self):
        self.assertEqual(parse_expected_text("[ROI1]\nShow [ROI2] on screen"), {1: "Show [ROI2] on screen"})


class BlockExcelTests(unittest.TestCase):
    def workbook_cases(self, headers, rows):
        with tempfile.TemporaryDirectory() as directory:
            book = Workbook()
            book.active.title = "tc_blocks"
            book.active.append(headers)
            for row in rows:
                book.active.append(row)
            path = Path(directory)/"cases.xlsx"
            book.save(path)
            book.close()
            before = path.read_bytes()
            cases, _ = load_cases(path)
            self.assertEqual(before, path.read_bytes())
            return cases

    def test_excel_multiline_blocks_resolve_original_ids_and_row_numbers(self):
        cases = self.workbook_cases(["Commands", "ROI", " Expected_Text "],
            [[None, "roi_courseop", "[ROI1]\nA\n[ROI2]\nB1\nB2\n[ROI3]\nC\n[ROI4]\nD"]])
        self.assertEqual(cases[0].expected, {1: "A", 2: "B1\nB2", 3: "C", 4: "D"})
        self.assertTrue(cases[0].expected_from_blocks)
        self.assertEqual(cases[0].row, 2)

    def test_empty_block_cell_keeps_legacy_columns_compatible(self):
        cases = self.workbook_cases(["Commands", "ROI", "Expected_Text", "Expected_1"], [[None, "p", None, "legacy"]])
        self.assertFalse(cases[0].expected_from_blocks)
        self.assertEqual(cases[0].expected, {1: "legacy"})

    def test_mixed_formats_are_rejected_even_when_ids_differ(self):
        with self.assertRaisesRegex(ValueError, "tc_blocks!2.*not both"):
            self.workbook_cases(["Commands", "ROI", "Expected_Text", "Expected_2"],
                                [[None, "p", "[ROI1]\na", "b"]])

    def test_duplicate_legacy_alias_ids_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.workbook_cases(["Commands", "ROI", "Expected_1", "Expected_ROI_1"], [[None, "p", "a", "b"]])

    def test_invalid_block_reports_sheet_and_row(self):
        with self.assertRaisesRegex(ValueError, "tc_blocks!2.*empty"):
            self.workbook_cases(["Commands", "ROI", "Expected_Text"], [[None, "p", "[ROI1]"]])

    def test_disabled_malformed_examples_are_skipped(self):
        cases = self.workbook_cases(["Commands", "ROI", "Expected_Text", "Execution"],
            [[None, "p", "not valid", "N"], [None, "p", "[ROI1]\nok", "Y"]])
        self.assertEqual(len(cases), 1)
        self.assertEqual(cases[0].row, 3)

    def test_block_ids_must_cover_every_saved_roi_without_image_or_preset_fallback(self):
        preset = {"image_size": [320, 240], "direction": "vertical", "scrolling": True,
                  "rois": [{"id": i, "rect": [0, 0, 20, 20], "expected": "preset fallback"} for i in (2, 5, 7)]}
        library = {"presets": {"p": preset}}
        case = TestCase("tc", 2, "test", "p", [], image="fallback.png", expected={2: "a", 5: "b"}, expected_from_blocks=True)
        before = copy.deepcopy(library)
        with self.assertRaisesRegex(ValueError, "missing ROI IDs.*7"):
            validate_cases([case], library, Path("."), require_anchors=False)
        with self.assertRaisesRegex(ValueError, "missing ROI IDs"):
            prepare_rois(None, case, preset, Path("."))
        case.expected[7] = "c"
        validate_cases([case], library, Path("."), require_anchors=False)
        case.expected[1] = "wrong ID"
        with self.assertRaisesRegex(ValueError, "unknown ROI IDs"):
            validate_cases([case], library, Path("."), require_anchors=False)
        self.assertEqual(library, before)


class TemplateTests(unittest.TestCase):
    def test_generated_xlsx_disabled_by_default_and_examples_parse_when_enabled(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/"template.xlsx"
            build_template(path)
            with self.assertRaisesRegex(ValueError, "No enabled test rows"):
                load_cases(path)
            book = load_workbook(path)
            self.assertIn("사용안내", book.sheetnames)
            sheet = book["tc_template"]
            self.assertEqual(sheet["G1"].value, "Expected_Text")
            self.assertTrue(sheet["G2"].alignment.wrap_text)
            self.assertTrue(sheet["G2"].value.startswith("[ROI1]\n"))
            for row in range(2, sheet.max_row+1):
                self.assertEqual(sheet.cell(row, 3).value, "N")
                sheet.cell(row, 3, "Y")
            book.save(path)
            book.close()
            cases, _ = load_cases(path)
            self.assertEqual([set(c.expected) for c in cases], [{1, 2, 3}, {1, 2, 3}, {1, 2, 3, 4}])
            self.assertIn("\n", cases[2].expected[1])
            self.assertTrue(all(c.expected_from_blocks for c in cases))


if __name__ == "__main__":
    unittest.main()