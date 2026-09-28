import copy
import json
from pathlib import Path
import queue
import tempfile
import threading
import unittest
from unittest.mock import MagicMock, patch
import zipfile

from openpyxl import Workbook
from PIL import Image

from ocr_roi_validator.automation import (
    Cancelled, TestCase, VestaSession, command_action, load_cases, number,
    product_config, product_config_names, reference_path, safe_extract,
)
from ocr_roi_validator.automation_ocr import prepare_rois, snapshot_ocr, validate_cases, verify_case
from ocr_roi_validator.automation_workflow import AutomationDialog
from ocr_roi_validator.gui import OCRValidatorGUI
from ocr_roi_validator.roi_presets import load_library, save_library
from ocr_roi_validator.window_anchor import (
    ClientWindow, choose_window, make_anchor, mapped_rect, selection_window, validate_anchor,
)


def preset(scrolling=False, direction="horizontal"):
    return {"image_size": [320, 240], "direction": direction, "scrolling": scrolling,
            "rois": [{"id": 1, "rect": [10, 20, 300, 60], "expected": "Hello"}],
            "capture_anchor": make_anchor(ClientWindow(123, "SDL_app", "Vesta", (100, 100, 420, 340)),
                                          (100, 100, 420, 340))}


class InlineCapture:
    def __init__(self, session, preset, fps, duration, stop):
        import time
        self.session, self.preset, self.stop = session, preset, stop
        self.end = time.monotonic()+duration
        self.count = 0
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass
    def next_frame(self, deadline):
        import time
        if time.monotonic() >= self.end:
            return None
        frame = self.session.frame(self.preset, None)
        self.count += 1
        return frame, time.monotonic()
    def stats(self):
        return {'captured_frames': self.count, 'dropped_frames': 0, 'pending_frames': 0}


def processor(text="Hello"):
    gui = OCRValidatorGUI.__new__(OCRValidatorGUI)
    settings = {"language_var": "en_es", "compare_mode_var": "exact", "similarity_threshold_var": "0.9",
                "vertical_list_mode_var": False, "scroll_mode_var": False, "roi_margin_var": "0",
                "min_roi_side_var": "32", "auto_upscale_var": False, "fast_long_roi_var": False,
                "context_detect_var": False, "context_margin_var": "0"}
    for key, value in settings.items():
        var = MagicMock()
        var.get.return_value = value
        setattr(gui, key, var)
    gui.engine = MagicMock()
    gui.engine.run.return_value = MagicMock(text=text, boxes=[object()], mean_score=0.99, n_boxes=1)
    value = snapshot_ocr(gui)
    value._capture_factory = InlineCapture
    return value


class GeometryTests(unittest.TestCase):
    def test_move_and_dpi_scale_including_negative_monitor(self):
        window = ClientWindow(1, "SDL", "GUI", (100, 200, 900, 800))
        anchor = make_anchor(window, (200, 350, 520, 590))
        self.assertEqual(mapped_rect(anchor, window.rect, (320, 240)), (200, 350, 520, 590))
        self.assertEqual(mapped_rect(anchor, (-1000, -300, 600, 900), (320, 240)), (-800, 0, -160, 480))

    def test_layout_change_rejected(self):
        with self.assertRaisesRegex(ValueError, "aspect"):
            mapped_rect(preset()["capture_anchor"], (0, 0, 640, 240), (320, 240))

    def test_invalid_anchors_rejected(self):
        for values in ([0, 0, 1.01, 1], [0, 0, float("nan"), 1], [0, 0, 0, 1], [0, 0, "1", 1]):
            anchor = preset()["capture_anchor"]
            anchor["normalized_rect"] = values
            with self.subTest(values=values), self.assertRaises(ValueError):
                validate_anchor(anchor)

    def test_selection_uses_smallest_containing_child(self):
        parent = ClientWindow(1, "P", "Parent", (0, 0, 800, 600))
        child = ClientWindow(2, "C", "Screen", (10, 10, 330, 250))
        self.assertEqual(selection_window([parent, child], (15, 15, 320, 240)), child)
        with self.assertRaises(ValueError):
            selection_window([parent], (-1, 0, 10, 20))

    def test_window_matching_rejects_ambiguity_and_supports_title_change(self):
        anchor = preset()["capture_anchor"]
        a = ClientWindow(1, "SDL_app", "New title", (0, 0, 320, 240))
        b = ClientWindow(2, "SDL_app", "Another title", a.rect)
        self.assertEqual(choose_window([a], anchor), a)
        with self.assertRaises(ValueError):
            choose_window([a, b], anchor)
        with self.assertRaises(ValueError):
            choose_window([], anchor)

    def test_anchor_json_roundtrip(self):
        data = {"version": 1, "presets": {"roi_courseop": preset()}}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roi.json"
            save_library(path, data)
            self.assertEqual(load_library(path), data)


class ExcelCommandTests(unittest.TestCase):
    def test_commands_tbl_sleep_and_unsupported_fail_closed(self):
        tables = {"go": {"Packet": "@0102", "Label": "UART1"}}
        self.assertEqual(command_action(":TBL go", tables), ("shell", "sim_uart write -l UART1 -d 0102"))
        self.assertEqual(command_action(':TBL {"Packet":"@20"}', {}), ("shell", "sim_irrc write -d 20"))
        self.assertEqual(command_action(":API System.sleep { time: 0.1 }", {}), ("sleep", 0.1))
        self.assertEqual(command_action("sim_key click power", {}), ("shell", "sim_key click power"))
        for command in (":TBL missing", ":LDB_PUSH x", ":API LMOSClient.set_config { data: a }", "reboot", ""):
            with self.subTest(command=command), self.assertRaises(ValueError):
                command_action(command, {})

    def test_workbook_multiple_sheets_headers_tbl_skips_and_line_breaks(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tc.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "tc_course"
            sheet.append(["Workbook title"])
            sheet.append([" No. ", " Commands ", "ROI", "Execution", "Capture_waiting_time", "Expected_1", "Image"])
            sheet.append([0, "sim_key power\n:TBL go\n:API System.sleep { time: 0 }", "roi_courseop", "Y", 0.25, "A\nB", "ref.png"])
            sheet.append([1, ":BAD", "", " n "])
            table = book.create_sheet("tbl_uart")
            table.append(["go", "uart", "UART1", "@0102"])
            second = book.create_sheet("tc_other")
            second.append(["Commands", "ROI"])
            second.append([None, "roi_courseop"])
            book.save(path)
            before = path.read_bytes()
            cases, tables = load_cases(path)
            self.assertEqual(len(cases), 2)
            self.assertEqual(cases[0].row, 3)
            self.assertEqual(cases[0].title, "0")
            self.assertEqual(cases[0].expected, {1: "A\nB"})
            self.assertEqual(cases[0].delay, 0.25)
            self.assertEqual(len(cases[0].commands), 3)
            self.assertEqual(tables["go"]["Packet"], "@0102")
            self.assertEqual(path.read_bytes(), before)

    def test_blank_roi_duplicate_headers_and_invalid_delay(self):
        for headers, row in ((["Commands", "ROI"], ["x", None]),
                             (["Commands", "ROI", "roi"], ["x", "p", "p"]),
                             (["Commands", "ROI", "Capture_waiting_time"], ["x", "p", -1])):
            with self.subTest(headers=headers), tempfile.TemporaryDirectory() as directory:
                book = Workbook()
                book.active.append(headers)
                book.active.append(row)
                path = Path(directory)/"bad.xlsx"
                book.save(path)
                with self.assertRaises(ValueError):
                    load_cases(path)

    def test_numeric_validation(self):
        for value in ("nan", "inf", -1):
            with self.assertRaises(ValueError):
                number(value, 1, "test")


class ArchiveSessionTests(unittest.TestCase):
    def test_actual_cfg_names_skip_global_and_target(self):
        names = ["vesta/sim_config/global.cfg", "vesta/sim_config/target.cfg",
                 "vesta\\sim_config\\washer.dqc.vented_combo_na.cfg", "vesta/sim_config/another.json"]
        self.assertEqual(product_config_names(names), ["another", "washer.dqc.vented_combo_na"])

    def test_safe_archive_and_product_detection(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root/"vesta.zip"
            with zipfile.ZipFile(archive, "w") as zipped:
                zipped.writestr("release/vesta.exe", b"test")
                zipped.writestr("release/sim_config/global.cfg", "{}")
                zipped.writestr("release/sim_config/target.cfg", "{}")
                zipped.writestr("release/sim_config/test.product.cfg", json.dumps({
                    "Port List": {"Socket": {"0": {"Client": "127.0.0.1:5888"}}}}))
            exe = safe_extract(archive, root/"unpacked", threading.Event())
            self.assertEqual(product_config(exe, ""), ("test.product", "127.0.0.1:5888"))

    def test_reject_zip_traversal_ads_symlink_and_case_collision(self):
        for name in ("../vesta.exe", "C:/vesta.exe", "folder/file:ads", "/vesta.exe", "folder./vesta.exe"):
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                path = Path(directory)/"bad.zip"
                with zipfile.ZipFile(path, "w") as zipped:
                    zipped.writestr(name, b"bad")
                with self.assertRaises(ValueError):
                    safe_extract(path, Path(directory)/"out", threading.Event())
                self.assertFalse((Path(directory)/"out").exists())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/"bad.zip"
            with zipfile.ZipFile(path, "w") as zipped:
                link = zipfile.ZipInfo("vesta.exe")
                link.external_attr = 0o120777 << 16
                zipped.writestr(link, "../target")
            with self.assertRaises(ValueError):
                safe_extract(path, Path(directory)/"out", threading.Event())

    def test_cancel_before_extraction(self):
        stop = threading.Event()
        stop.set()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/"test.zip"
            with zipfile.ZipFile(path, "w") as zipped:
                zipped.writestr("vesta.exe", b"x")
            with self.assertRaises(Cancelled):
                safe_extract(path, Path(directory)/"out", stop)

    def test_sequential_commands_delay_and_cancel(self):
        session = VestaSession(Path("lite"), Path("output"))
        session.shell = MagicMock()
        case = TestCase("tc", 2, "test", "p", ["first", "second"])
        session.execute(case, {}, threading.Event())
        self.assertEqual([c.args[0] for c in session.shell.call_args_list], ["first", "second"])
        stop = threading.Event()
        stop.set()
        with self.assertRaises(Cancelled):
            session.execute(case, {}, stop)

    def test_command_gaps_and_capture_wait_are_separate(self):
        session = VestaSession(Path("lite"), Path("output"))
        events = []
        session.shell = lambda command, stop: events.append(command)
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda seconds: events.append(seconds) or False
        session.execute(TestCase("tc", 2, "test", "p", ["A", "B", "C"], delay=3), {}, stop)
        self.assertEqual(events, ["A", 1.0, "B", 1.0, "C", 3])

    def test_explicit_sleep_replaces_default_gap(self):
        session = VestaSession(Path("lite"), Path("output"))
        events = []
        session.shell = lambda command, stop: events.append(command)
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda seconds: events.append(seconds) or False
        commands = ["A", ":API System.sleep { time: 2 }", "B", "C"]
        session.execute(TestCase("tc", 2, "test", "p", commands), {}, stop)
        self.assertEqual(events, ["A", 2.0, "B", 1.0, "C", 0])

    def test_cancel_during_gap_does_not_send_next_command(self):
        session = VestaSession(Path("lite"), Path("output"))
        session.shell = MagicMock()
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.return_value = True
        with self.assertRaises(Cancelled):
            session.execute(TestCase("tc", 2, "test", "p", ["A", "B"]), {}, stop)
        self.assertEqual([c.args[0] for c in session.shell.call_args_list], ["A"])

    def test_reference_lookup_rejects_traversal_and_ambiguous_names(self):
        case = TestCase("tc", 2, "t", "p", [], image="../outside.png")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(ValueError):
                reference_path(case, root)
            case.image = "answer.png"
            for folder in ("a", "b"):
                (root/folder).mkdir()
                Image.new("RGB", (1, 1)).save(root/folder/case.image)
            with self.assertRaises(ValueError):
                reference_path(case, root)


class VerificationTests(unittest.TestCase):
    def test_missing_expectations_and_missing_anchors_fail_preflight(self):
        case = TestCase("tc", 2, "t", "roi_courseop", [])
        data = {"version": 1, "presets": {"roi_courseop": preset()}}
        validate_cases([case], data, Path("."))
        del data["presets"][case.preset]["capture_anchor"]
        with self.assertRaisesRegex(ValueError, "register"):
            validate_cases([case], data, Path("."))
        data["presets"][case.preset] = preset()
        data["presets"][case.preset]["rois"][0]["expected"] = ""
        with self.assertRaisesRegex(ValueError, "empty"):
            validate_cases([case], data, Path("."))

    def test_expected_override_and_unknown_roi(self):
        case = TestCase("tc", 2, "t", "roi_courseop", [], expected={1: "Override"})
        p = preset()
        self.assertEqual(prepare_rois(processor(), case, p, Path("."))[1].expected, "Override")
        self.assertEqual(p["rois"][0]["expected"], "Hello")
        case.expected[2] = "invalid"
        with self.assertRaisesRegex(ValueError, "unknown ROI IDs"):
            validate_cases([case], {"presets": {case.preset: p}}, Path("."))

    def test_reference_ocr_only_fills_missing_and_never_mutates_preset(self):
        p = preset()
        p["rois"][0]["expected"] = ""
        before = copy.deepcopy(p)
        case = TestCase("tc", 2, "t", "p", [], image="expected.png")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new("RGB", (320, 240)).save(root/case.image)
            prepared = prepare_rois(processor("Ground truth"), case, p, root)
            self.assertEqual(prepared[1].expected, "Ground truth")
        self.assertEqual(p, before)

    def test_snapshot_vars_do_not_access_tk_in_worker(self):
        p = processor()
        self.assertEqual(p.language_var.get(), "en_es")
        self.assertEqual(type(p.language_var).__name__, "FrozenValue")

    @patch("ocr_roi_validator.automation_ocr.mss.mss")
    def test_static_pass_and_fail_with_exact_diagnostics(self, capture):
        for text, status in (("Hello", "PASS"), ("Wrong", "FAIL_TIMEOUT")):
            with self.subTest(text=text), tempfile.TemporaryDirectory() as directory:
                session = MagicMock()
                session.frame.return_value = Image.new("RGB", (320, 240))
                result, frame = verify_case(processor(text), session, TestCase("tc", 2, "t", "p", []),
                    preset(), Path(directory), 0.8, 20, threading.Event(), Path(directory)/"result")
                self.assertEqual(result["status"], status)
                self.assertGreaterEqual(session.frame.call_count, 2)
                if status != "PASS":
                    metadata = json.loads((Path(result["rois"][0]["diagnostic"])/"metadata.json").read_text())
                    self.assertEqual(metadata["ocr_input_fidelity"], "exact_recorded_ocr_input")

    @patch("ocr_roi_validator.automation_ocr.mss.mss")
    def test_horizontal_and_vertical_reject_partial_content(self, capture):
        for direction in ("horizontal", "vertical"):
            with self.subTest(direction=direction), tempfile.TemporaryDirectory() as directory:
                session = MagicMock()
                session.frame.return_value = Image.new("RGB", (320, 240))
                p = preset(True, direction)
                p["rois"][0]["expected"] = "Hello\nWorld\nFinal option" if direction == "vertical" else "Hello world with many more options"
                result, _ = verify_case(processor(), session, TestCase("tc", 2, "t", "p", []),
                    p, Path(directory), 0.00001, 8, threading.Event(), Path(directory)/"result")
                self.assertEqual(result["status"], "FAIL_TIMEOUT")

    @patch("ocr_roi_validator.automation_ocr.mss.mss")
    def test_content_policy_ends_after_confirmation_before_maximum(self, capture):
        clock = [0.0]
        session = MagicMock()
        def next_frame(*args):
            clock[0] += 0.3
            return Image.new("RGB", (320, 240))
        session.frame.side_effect = next_frame
        stop = MagicMock()
        stop.is_set.return_value = False
        with tempfile.TemporaryDirectory() as directory, patch(
                "ocr_roi_validator.automation_ocr.time.monotonic", side_effect=lambda: clock[0]):
            result, _ = verify_case(processor(), session, TestCase("tc", 2, "t", "p", []),
                preset(True), Path(directory), 30, 100, stop, Path(directory)/"result")
        self.assertEqual(result["status"], "PASS")
        self.assertEqual(session.frame.call_count, 3)

    @patch("ocr_roi_validator.automation_ocr.mss.mss")
    def test_cancel_never_returns_pass(self, capture):
        stop = threading.Event()
        stop.set()
        with tempfile.TemporaryDirectory() as directory:
            result, _ = verify_case(processor(), MagicMock(), TestCase("tc", 2, "t", "p", []), preset(),
                        Path(directory), 1, 1, stop, Path(directory)/"result")
            self.assertEqual(result["status"], "CANCELLED")


class ReportWorkflowTests(unittest.TestCase):
    def run_report(self, directory, error=None):
        dialog = AutomationDialog.__new__(AutomationDialog)
        dialog.busy = lambda: False
        dialog.output = Path(directory)
        dialog.stop = threading.Event()
        dialog.events = queue.Queue()
        dialog.tree = MagicMock()
        dialog.gui = MagicMock()
        dialog.gui.preset_library = {"version": 1, "presets": {"p": preset()}}
        dialog.ensure_session = MagicMock()
        session = MagicMock()
        dialog.session = session
        dialog.start_worker = lambda task: task()
        dialog.config = lambda: {"excel": str(Path(directory)/"tc.xlsx"), "zip": "vesta.zip",
                                 "references": directory, "roi_column": "ROI", "duration": 1, "fps": 1}
        cases = [TestCase("tc", i+2, f"Case {i}", "p", [f"command{i}"]) for i in range(2)]
        p = processor()
        session.execute.side_effect = error
        with patch("ocr_roi_validator.automation_workflow.load_cases", return_value=(cases, {})), \
             patch("ocr_roi_validator.automation_workflow.snapshot_ocr", return_value=p), \
             patch("ocr_roi_validator.automation_workflow.verify_case", side_effect=lambda *args: (
                 {"status": "PASS", "rois": []}, Image.new("RGB", (320, 240)))):
            if isinstance(error, Cancelled):
                with self.assertRaises(type(error)):
                    dialog.run()
            else:
                dialog.run()
        report = json.loads(next(Path(directory).glob("results_*.json")).read_text())
        session.close.assert_called_once()
        self.assertIsNone(dialog.session)
        return report, session

    def test_success_records_original_sheet_rows_and_closes_session(self):
        with tempfile.TemporaryDirectory() as directory:
            report, session = self.run_report(directory)
            self.assertEqual(report["state"], "COMPLETED")
            self.assertEqual([r["row"] for r in report["results"]], [2, 3])
            self.assertEqual(session.execute.call_count, 2)

    def test_command_error_continues_independent_rows_and_records_error(self):
        with tempfile.TemporaryDirectory() as directory:
            report, session = self.run_report(directory, RuntimeError("transport failed"))
            self.assertEqual(report["state"], "COMPLETED")
            self.assertEqual(report["results"][0]["status"], "ERROR")
            self.assertEqual(session.execute.call_count, 2)

    def test_cancel_records_cancelled_state_never_pass(self):
        with tempfile.TemporaryDirectory() as directory:
            report, session = self.run_report(directory, Cancelled("stop"))
            self.assertEqual(report["state"], "CANCELLED")
            self.assertEqual([r["status"] for r in report["results"]], ["CANCELLED", "NOT_RUN"])


if __name__ == "__main__":
    unittest.main()