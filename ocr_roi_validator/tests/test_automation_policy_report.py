import copy
import multiprocessing as mp
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from PIL import Image

from ocr_roi_validator.automation import Cancelled, TestCase, verification_options
from ocr_roi_validator.automation_ocr import verify_case
from ocr_roi_validator.automation_report import initial_result, write_excel_report
from ocr_roi_validator.isolated_ocr import IsolatedOCR
from ocr_roi_validator.verification_policy import CycleTracker
import test_automation as helpers
from test_automation import preset, processor


def blocked_ocr(connection):
    connection.recv()
    # Deliberately hang as an unresponsive native OCR call would.
    threading.Event().wait(60)


class CycleTests(unittest.TestCase):
    def test_complete_cycles_are_counted_independently(self):
        tracker = CycleTracker("ABCDEFGHIJKLMNOP")
        for text in ["GHIJ", "JKLM", "MNOP", "ABCD"]:
            tracker.add(text)
        self.assertEqual(tracker.cycles, 0)  # first partial traversal is not a cycle
        for _ in range(2):
            for text in ["ABCD", "DEFG", "GHIJ", "JKLM", "MNOP", "ABCD"]:
                tracker.add(text)
        self.assertEqual(tracker.cycles, 2)

    def test_repeated_word_is_ambiguous_not_a_wrap(self):
        tracker = CycleTracker("ONE TWO ONE END")
        for text in ["ONE", "ONE", "ONE", "ONE"]:
            tracker.add(text)
        self.assertEqual(tracker.cycles, 0)
        self.assertEqual(tracker.ambiguous, 4)

    def test_static_full_text_never_counts_as_cycle(self):
        tracker = CycleTracker("Hello world")
        for _ in range(20):
            tracker.add("Hello world")
        self.assertEqual(tracker.cycles, 0)

    def test_gap_or_backward_order_does_not_count(self):
        for frames in (["ABCD", "MNOP", "ABCD"], ["ABCD", "DEFG", "ABCD"]):
            tracker = CycleTracker("ABCDEFGHIJKLMNOP")
            for frame in frames:
                tracker.add(frame)
            self.assertEqual(tracker.cycles, 0)

    def test_options_enforce_thirty_second_cap_and_integer_cycles(self):
        self.assertEqual(verification_options({})[:3], (30, "content", 1))
        self.assertEqual(verification_options({"verificationmode": "cycles", "requiredcycles": 2,
                                              "verificationmode1": "content"})[3], {1: "content"})
        for values in ({"maxobservationsec": 31}, {"maxobservationsec": 0}, {"requiredcycles": 1.5},
                       {"requiredcycles": -1}, {"verificationmode": "guess"}, {"verificationmode1": "guess"}):
            with self.subTest(values=values), self.assertRaises(ValueError):
                verification_options(values)


class MultiROITests(unittest.TestCase):
    def run_rois(self, mode="content", cancel=False, third_never_passes=False):
        p = preset()
        p["rois"] = [{"id": i, "rect": [i, 0, i+10, 20], "expected": f"Text {i}"} for i in (1, 2, 3)]
        clock = [0.0]
        frames = [0]
        calls = {1: 0, 2: 0, 3: 0}
        proc = processor()
        stop = threading.Event()
        session = MagicMock()

        def frame(*args):
            frames[0] += 1
            clock[0] += 0.3
            if cancel and frames[0] == 5:
                stop.set()
            return Image.new("RGB", (320, 240))

        def ocr(image, rect, expected):
            roi_id = rect[0]
            calls[roi_id] += 1
            text = expected if roi_id != 3 or (frames[0] >= 4 and not third_never_passes) else "Wrong"
            return MagicMock(text=text, mean_score=0.99)

        proc._run_roi_ocr = ocr
        session.frame.side_effect = frame
        with tempfile.TemporaryDirectory() as directory, \
             patch("ocr_roi_validator.automation_ocr.time.monotonic", side_effect=lambda: clock[0]), \
             patch("ocr_roi_validator.automation_ocr.mss.mss"):
            case = TestCase("tc", 2, "test", "p", [], verification_mode=mode)
            result, _ = verify_case(proc, session, case, p, Path(directory), 2.5, 1000, stop, Path(directory)/"result")
        return result, calls

    def test_only_unfinished_roi_continues_and_all_pass_ends_early(self):
        result, calls = self.run_rois()
        self.assertEqual(result["status"], "PASS")
        self.assertEqual(calls[1], 3)
        self.assertEqual(calls[2], 3)
        self.assertGreater(calls[3], calls[1])
        self.assertLess(result["elapsed_sec"], 2.5)

    def test_timeout_preserves_completed_rois(self):
        result, _ = self.run_rois(third_never_passes=True)
        self.assertEqual(result["status"], "FAIL_TIMEOUT")
        self.assertEqual([r["status"] for r in result["rois"]], ["PASS", "PASS", "FAIL_TIMEOUT"])

    def test_cancel_preserves_completed_rois_and_partial_actual(self):
        result, _ = self.run_rois(cancel=True)
        self.assertEqual(result["status"], "CANCELLED")
        self.assertEqual([r["status"] for r in result["rois"]], ["PASS", "PASS", "CANCELLED"])
        self.assertTrue(result["rois"][2]["actual"])

    def test_cycle_mode_never_accepts_static_full_text(self):
        result, _ = self.run_rois(mode="cycles")
        self.assertEqual(result["status"], "FAIL_TIMEOUT")


class WatchdogTests(unittest.TestCase):
    def make_service(self):
        context = mp.get_context("spawn")
        parent, child = context.Pipe()
        service = IsolatedOCR(MagicMock())
        service.connection = parent
        service.process = context.Process(target=blocked_ocr, args=(child,), daemon=True)
        service.process.start()
        child.close()
        return service

    def test_hung_ocr_is_terminated_at_deadline(self):
        service = self.make_service()
        started = time.monotonic()
        try:
            with self.assertRaises(TimeoutError):
                service.run(Image.new("RGB", (10, 10)), (0, 0, 10, 10), "test", threading.Event(), started+0.2)
            self.assertIsNone(service.process)
            self.assertLess(time.monotonic()-started, 3)
        finally:
            service.close()

    def test_cancel_terminates_ocr(self):
        service = self.make_service()
        stop = threading.Event()
        stop.set()
        try:
            with self.assertRaises(Cancelled):
                service.run(Image.new("RGB", (10, 10)), (0, 0, 10, 10), "test", stop, time.monotonic()+30)
            self.assertIsNone(service.process)
        finally:
            service.close()


class ExcelReportTests(unittest.TestCase):
    def test_report_preserves_fields_and_never_executes_formulas(self):
        case = TestCase("tc_sheet", 9, "테스트", "roi_courseop", ["=HYPERLINK(\"bad\")", "second"],
                        image="answer.png", expected={1: "=1+1\n다음 줄"}, tc_number="TC-009")
        entry = initial_result(case, preset())
        entry.update(status="CANCELLED")
        entry["rois"][0].update(status="CANCELLED", actual="=2+2\x01\n검출")
        before = copy.deepcopy(entry)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/"report.xlsx"
            write_excel_report(path, {"state": "CANCELLED", "results": [entry]})
            book = load_workbook(path, data_only=False)
            self.assertEqual(book.sheetnames, ["TC Summary", "ROI Results", "OCR Timing", "Run Info"])
            row = list(book["ROI Results"].iter_rows(min_row=2))[0]
            self.assertEqual(row[0].value, "TC-009")
            self.assertEqual(row[4].value, "answer.png")
            self.assertEqual(row[5].data_type, "s")
            self.assertEqual(row[6].value, "roi_courseop")
            self.assertEqual(row[8].value, "=1+1\n다음 줄")
            self.assertEqual(row[8].data_type, "s")
            self.assertEqual(row[9].value, "=2+2\n검출")
            book.close()
        self.assertEqual(entry, before)

    def test_completion_and_cancel_export_automatically(self):
        helper = helpers.ReportWorkflowTests()
        for error in (None, Cancelled("stop")):
            with self.subTest(error=error), tempfile.TemporaryDirectory() as directory:
                report, _ = helper.run_report(directory, error)
                path = next(Path(directory).glob("results_*.xlsx"))
                book = load_workbook(path)
                self.assertEqual(book["TC Summary"].max_row, 3)
                if error:
                    self.assertEqual(book["ROI Results"]["K2"].value, "CANCELLED")
                    self.assertEqual(book["ROI Results"]["K3"].value, "NOT_RUN")
                book.close()


if __name__ == "__main__":
    unittest.main()