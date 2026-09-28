import copy
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
from PIL import Image
from openpyxl import load_workbook
from ocr_roi_validator.automation import TestCase
from ocr_roi_validator.automation_ocr import verify_case
from ocr_roi_validator.automation_report import write_excel_report
from test_automation import processor, preset
from test_static_and_diacritics import result


class SynchronousCaptureTests(unittest.TestCase):
    def test_capture_waits_for_all_rois_and_records_every_attempt(self):
        proc, saved = processor(), preset(True)
        second = copy.deepcopy(saved['rois'][0])
        second['id'] = 2
        second['rect'] = [10, 70, 300, 110]
        saved['rois'].append(second)
        clock, events = [0.0], []
        def frame(*args):
            events.append('capture')
            clock[0] += .2
            return Image.new('RGB', (320, 240))
        def ocr(*args):
            events.append('ocr')
            clock[0] += .05
            return result('Wrong')
        session = MagicMock()
        session.frame.side_effect = frame
        proc._run_roi_ocr = ocr
        with tempfile.TemporaryDirectory() as folder, patch('ocr_roi_validator.automation_ocr.mss.mss'), patch('ocr_roi_validator.automation_ocr.time.monotonic', side_effect=lambda: clock[0]):
            report, _ = verify_case(proc, session, TestCase('tc', 2, 'test', 'p', []), saved,
                Path(folder), .95, 1000, threading.Event(), Path(folder)/'out')
            self.assertEqual(events[:6], ['capture', 'ocr', 'ocr', 'capture', 'ocr', 'ocr'])
            detail = report['rois'][0]['details']
            self.assertEqual(detail['capture']['capture_policy'], 'synchronous_de22028')
            self.assertEqual(detail['ocr_calls'], len(detail['frame_timings']))
            self.assertGreater(detail['ocr_calls'], len(detail['raw_observations']))
            for timing in detail['frame_timings']:
                self.assertLessEqual(timing['captured_sec'], timing['ocr_started_sec'])
                self.assertLess(timing['ocr_started_sec'], timing['ocr_finished_sec'])
                self.assertEqual(timing['raw_text'], 'Wrong')
            tc = dict(report, sheet='tc', row=2, title='test', preset='p')
            path = Path(folder)/'report.xlsx'
            write_excel_report(path, {'results': [tc]})
            book = load_workbook(path)
            self.assertEqual(book['OCR Timing'].max_row, 1+sum(r['details']['ocr_calls'] for r in report['rois']))
            self.assertTrue(book['ROI Results'].column_dimensions['V'].hidden)
            self.assertTrue(book['ROI Results'].column_dimensions['S'].hidden)
            self.assertFalse(book['ROI Results'].column_dimensions['J'].hidden)
            self.assertIsNotNone(book['ROI Results']['J2'].comment)
            book.close()
