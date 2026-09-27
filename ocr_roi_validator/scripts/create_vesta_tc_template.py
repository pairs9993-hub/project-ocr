"""Generate the editable, macro-free Excel template for Vesta ROI automation."""
from pathlib import Path
import argparse

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


def build_template(output: Path) -> None:
    book = Workbook()
    guide = book.active
    guide.title = "사용안내"
    guide.append(["항목", "설명"])
    instructions = [
        ("1. 양식", "tc_template 시트의 예시를 실제 TC로 수정하세요. 예시는 모두 Execution=N이며 실행되지 않습니다. 준비된 행만 Y로 변경하세요."),
        ("2. 준비", "run.bat → 자동화 검증 → Vesta ZIP과 이 Excel 첨부. 기존 ROI는 최초 1회 Vesta 화면 기준 등록이 필요합니다."),
        ("3. ROI", "ROI 열에는 저장된 프리셋 이름(예: roi_courseop)을 입력합니다. ROI ID는 화면 위아래 순서가 아니라 저장된 번호입니다."),
        ("4. Expected_Text", "한 셀에 [ROI1] 줄바꿈 정답 줄바꿈 [ROI2] 줄바꿈 정답… 형식으로 입력합니다. 구분자는 반드시 독립된 줄입니다."),
        ("5. 줄바꿈", "Excel에서 Alt+Enter. 다음 [ROI번호] 전까지 같은 ROI 정답입니다. 정답 내부 줄바꿈·공백은 보존하고 블록 앞뒤 빈 줄만 제거합니다."),
        ("6. 필수 정답", "Expected_Text를 사용하면 프리셋의 모든 ROI를 작성해야 합니다. 빠진/중복/없는 ID, 빈 정답은 오류이며 프리셋이나 Image로 대신 채우지 않습니다."),
        ("7. 기존 열 방식", "Expected_1, Expected_2… 개별 열도 지원합니다. 동일 행에서 Expected_Text와 개별 정답 열을 동시에 채우면 오류입니다."),
        ("8. Commands", "실제 제품용 LITE 명령을 입력하세요. 여러 명령은 Alt+Enter로 구분합니다. 빈 셀은 명령 전송 없이 현재 화면 검증입니다. 예시 명령은 대기만 수행합니다."),
        ("9. 지원 명령", "일반 LDB shell / :TBL / :API System.sleep. LMOS/LCP API, reboot, LDB_INSTALL 등은 현재 지원하지 않습니다."),
        ("10. Capture_waiting_time", "모든 명령 후 OCR 캡처를 시작하기 전에 기다리는 시간(초)입니다. 스크롤 검증 시간과 다릅니다."),
        ("11. ROI 관찰 종료", "모든 ROI가 PASS면 조기 종료합니다. 최대 시간은 GUI와 Max_observation_sec 중 작은 값이며 상한 30초입니다. PASS인 ROI는 고정하고 나머지만 관찰합니다."),
        ("12. 독립 TC 진행", "PASS/FAIL_TIMEOUT/ERROR는 결과 저장 후 다음 TC로 진행합니다. 사용자 중지 시 전체 실행을 중단하고 미실행 TC를 구분해 Excel로 자동 저장합니다."),
        ("13. Image (선택)", "Expected_Text를 비워 둔 기존 방식에서만 부족한 정답을 이미지 OCR로 채울 수 있습니다. 정답 폴더 기준 경로/파일명, 프리셋과 동일 화면 크기가 필요합니다."),
        ("14. 방향", "Horizontal/Vertical은 프리셋 설정을 따릅니다. 셀에 여러 줄을 입력해도 방향이 자동 변경되지는 않습니다."),
        ("15. 예시 주의", "예시 문구·명령·roi_example_vertical_4는 설명용입니다. 실제 제품 문구/명령으로 교체하고 해당 프리셋을 저장한 뒤 실행하세요."),
        ("16. 검증 방식", "Verification_mode=content는 내용 확인, cycles는 Required_cycles회 정상 순환을 요구합니다. 개별 ROI는 Verification_mode_1 등의 열로 덮어쓸 수 있습니다. 정적인 ROI는 content 사용."),
        ("17. 주기 한계", "고유한 문자열 구간이 정답 순서로 겹쳐 이동하고 전체 내용 후 시작점에 돌아와야 1회입니다. 반복 단어가 모호하거나 OCR 누락/오류로 확인 불가하면 횟수를 추측하지 않고 시간초과 처리합니다."),
        ("18. 결과", "완료/중지 시 TC Summary, ROI Results, Run Info 시트가 자동 생성됩니다. 정답지 이름(Image)은 사용 여부와 관계없이 결과에 남습니다. 성공 확인은 0.5초 이상, 최소 2회 관측입니다."),
    ]
    for row in instructions:
        guide.append(row)
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 110
    for row in range(2, guide.max_row+1):
        guide.row_dimensions[row].height = 50
    guide.freeze_panes = "B2"

    sheet = book.create_sheet("tc_template")
    headers = ["No.", "Title", "Execution", "Commands", "Capture_waiting_time", "ROI", "Expected_Text", "Image", "Note",
               "Max_observation_sec", "Verification_mode", "Required_cycles", "Verification_mode_1"]
    sheet.append(headers)
    sheet.append([1, "코스 옵션 — 3 ROI", "N", ":API System.sleep { time: 1 }", 2,
                  "roi_courseop", "[ROI1]\n첫 번째 ROI의 전체 정답 문구\n\n[ROI2]\n두 번째 ROI의 전체 정답 문구\n\n[ROI3]\n세 번째 ROI의 전체 정답 문구",
                  None, "예시입니다. 실제 정답/제품 명령으로 바꾼 뒤 Y로 변경하세요. 명령은 대기만 수행합니다."])
    sheet.append([2, "코스 옵션 — 다른 정답", "N", None, 2, "roi_courseop",
                  "[ROI1]\n첫 번째 영역의 변경된 정답\n\n[ROI2]\n두 번째 영역의 변경된 정답\n\n[ROI3]\n세 번째 영역의 변경된 정답",
                  None, "같은 프리셋도 TC마다 다른 정답을 입력할 수 있습니다. Commands가 비면 현재 화면을 검증합니다."])
    sheet.append([3, "Vertical 여러 줄 / 4 ROI 예시", "N", None, 2, "roi_example_vertical_4",
                  "[ROI1]\n첫 번째 항목\n두 번째 항목\n세 번째 항목\n\n[ROI2]\n옵션 A\n옵션 B\n\n[ROI3]\n안내 첫째 줄\n안내 둘째 줄\n\n[ROI4]\n네 번째 ROI 정답",
                  None, "이 프리셋은 기본 제공되지 않습니다. ROI 4개와 Vertical 설정을 가진 실제 프리셋 이름으로 바꾸세요."])
    for row in range(2, 5):
        sheet.cell(row, 10, 30)
        sheet.cell(row, 11, "cycles" if row == 3 else "content")
        sheet.cell(row, 12, 2 if row == 3 else 1)
    sheet.cell(3, 13, "content")
    widths = {"A": 8, "B": 30, "C": 13, "D": 42, "E": 25, "F": 29, "G": 64, "H": 30, "I": 65,
              "J": 25, "K": 25, "L": 22, "M": 28}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in (2, 3):
        sheet.row_dimensions[row].height = 185
    sheet.row_dimensions[4].height = 290
    sheet.freeze_panes = "G2"
    table = Table(displayName="VestaTestCases", ref=f"A1:M{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    execution = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    execution.errorTitle = "Execution"
    execution.error = "Y 또는 N을 선택하세요."
    execution.showErrorMessage = True
    execution.errorStyle = "stop"
    sheet.add_data_validation(execution)
    execution.add("C2:C10000")
    delay = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
    delay.error = "0 이상의 초 단위 숫자를 입력하세요."
    delay.showErrorMessage = True
    delay.errorStyle = "stop"
    sheet.add_data_validation(delay)
    delay.add("E2:E10000")
    sheet["G1"].comment = Comment("[ROI1]을 독립된 줄로 쓰고 다음 줄부터 정답을 입력하세요. Alt+Enter로 줄바꿈. 모든 ROI 필수. 내부 줄바꿈 보존.", "ROI Validator")
    sheet["E1"].comment = Comment("명령 후 화면 전환 대기 시간입니다. 스크롤 검증 시간은 자동화 창에서 별도로 설정합니다.", "ROI Validator")
    sheet["C1"].comment = Comment("안전을 위해 예시 행은 모두 N입니다. 실제 TC를 작성한 뒤 Y로 바꾸세요.", "ROI Validator")
    for page in book:
        page.sheet_view.showGridLines = False
        page.sheet_view.zoomScale = 85
        for row in page:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="맑은 고딕", size=11)
        for cell in page[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=11)
        page.row_dimensions[1].height = 30
    book.active = 1
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "examples" / "vesta_automation_tc_template.xlsx")
    args = parser.parse_args()
    if args.output.exists():
        parser.error(f"Refusing to overwrite an existing workbook: {args.output}")
    build_template(args.output)
    print(args.output)