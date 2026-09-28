"""Portable, versioned ROI libraries and read-only Excel references."""
from __future__ import annotations

import json
from pathlib import Path


def validate_library(data: dict) -> dict:
    if not isinstance(data, dict) or data.get("version") != 1:
        raise ValueError("ROI library must have version: 1.")
    presets = data.get("presets")
    if not isinstance(presets, dict):
        raise ValueError("ROI library must contain a presets object.")
    if "shared_capture" in data:
        shared = data["shared_capture"]
        if not isinstance(shared, dict):
            raise ValueError("Invalid shared capture reference.")
        size = shared.get("image_size")
        if not isinstance(size, list) or len(size) != 2 or any(type(v) is not int or v <= 0 for v in size):
            raise ValueError("Invalid shared capture image_size.")
        from .window_anchor import validate_anchor
        validate_anchor(shared.get("capture_anchor"))
    for name, preset in presets.items():
        if not isinstance(name, str) or not name.strip() or name != name.strip():
            raise ValueError("Preset names must be nonempty without surrounding spaces.")
        if not isinstance(preset, dict):
            raise ValueError(f"{name}: invalid preset.")
        if "capture_anchor" in preset:
            from .window_anchor import validate_anchor
            validate_anchor(preset["capture_anchor"])
        if preset.get("direction") not in ("horizontal", "vertical"):
            raise ValueError(f"{name}: direction must be horizontal or vertical.")
        if type(preset.get("scrolling")) is not bool:
            raise ValueError(f"{name}: scrolling must be a boolean.")
        size = preset.get("image_size")
        if not isinstance(size, list) or len(size) != 2 or any(type(v) is not int or v <= 0 for v in size):
            raise ValueError(f"{name}: invalid image_size.")
        rois = preset.get("rois")
        if not isinstance(rois, list) or not rois:
            raise ValueError(f"{name}: at least one ROI is required.")
        ids = set()
        for roi in rois:
            if not isinstance(roi, dict):
                raise ValueError(f"{name}: invalid ROI.")
            roi_id, rect = roi.get("id"), roi.get("rect")
            if type(roi_id) is not int or roi_id <= 0 or roi_id in ids:
                raise ValueError(f"{name}: ROI IDs must be unique positive integers.")
            ids.add(roi_id)
            if not isinstance(rect, list) or len(rect) != 4 or any(type(v) is not int for v in rect):
                raise ValueError(f"{name}: invalid ROI rectangle.")
            x1, y1, x2, y2 = rect
            if not (0 <= x1 < x2 <= size[0] and 0 <= y1 < y2 <= size[1]):
                raise ValueError(f"{name}: ROI {roi_id} is outside image_size.")
            if not isinstance(roi.get("expected", ""), str):
                raise ValueError(f"{name}: expected text must be a string.")
    return data


def load_library(path: Path) -> dict:
    return validate_library(json.loads(path.read_text(encoding="utf-8-sig")))


def save_library(path: Path, data: dict) -> None:
    validate_library(data)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def resolve_preset(library: dict, name: str, image_size: tuple[int, int]) -> dict:
    name = name.strip()
    if not name:
        raise ValueError("The selected Excel row has no ROI preset name.")
    if name not in library["presets"]:
        raise ValueError(f"Unknown ROI preset: {name}. Import its JSON library first.")
    preset = library["presets"][name]
    if tuple(preset["image_size"]) != image_size:
        raise ValueError(f"{name}: expected image/capture size {preset['image_size']}, got {list(image_size)}. Select a matching capture area.")
    return preset


def read_excel(path: Path) -> dict[str, list[tuple]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def excel_references(rows: list[tuple], column: int) -> list[tuple[int, str]]:
    """First row is the header; retain original Excel row numbers."""
    return [(number, str(row[column]).strip() if column < len(row) and row[column] is not None else "")
            for number, row in enumerate(rows[1:], start=2)]


def automation_preset(library: dict, name: str) -> dict:
    """Capture the shared screen, then map it to this preset's own ROI coordinates."""
    from copy import deepcopy
    preset = deepcopy(library["presets"][name])
    shared = library.get("shared_capture")
    if shared:
        preset["capture_anchor"] = deepcopy(shared["capture_anchor"])
        preset["capture_size"] = list(shared["image_size"])
    return preset


def with_shared_capture(library: dict, name: str, anchor: dict, image_size=None) -> dict:
    from copy import deepcopy
    updated = deepcopy(library)
    # Legacy callers default to the selected preset's size; a new calibration
    # records the actual selected screen size, independently of any preset.
    selected_size = image_size if image_size is not None else library["presets"][name]["image_size"]
    updated["shared_capture"] = {"image_size": list(selected_size),
                                 "capture_anchor": deepcopy(anchor)}
    return validate_library(updated)
